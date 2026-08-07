"""Build one master Excel workbook consolidating, per headline-eligible
question: the schema/gold info, every annotator's raw commitment judgement,
and every debate model's final answer — for side-by-side inspection.

Reanalysis only; no model calls; nothing overwritten. Output:
    gold_commitment_master_comparison.xlsx  (repo root)

Sheets:
    gold_and_annotators  — gold + schema + fable/codex/antigravity/human calls
    model_predictions    — each model's final-round predicted label vs gold
    commitment_disagreements — the rows where annotators disagree on commitment
    legend               — column definitions
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO))

from src.schema import load_schemas  # noqa: E402

ANA = REPO / "analysis"
RES = REPO / "results"
OUT = REPO / "gold_commitment_master_comparison.xlsx"
GOLD_CSV = Path(r"C:\Users\ddkxi\PycharmProjects\Contrastive learning"
                r"\NASDAQ processed data\Final_benchmark\final_dataset_release"
                r"\FinTradeBench_Golden_Seed_150.csv")
schemas = load_schemas()
ELIG = [s for s in schemas.values() if s.headline_eligible]


def schema_commit(q):
    return "noncommitted" if schemas[q].gold_label in schemas[q].noncommit_set else "committed"


# full gold final-answer text
gold_text = {}
if GOLD_CSV.exists():
    src = pd.read_csv(GOLD_CSV)
    for _, r in src.iterrows():
        resp = str(r["response"])
        parts = re.split(r"(?i)final answer\s*:?", resp)
        tail = parts[-1].strip() if len(parts) > 1 else resp[-2500:]
        gold_text[str(r["question_id"])] = re.sub(r"\s+", " ", tail).strip()[:2500]


def load_annot(path, lcol, ccol, confcol, notecol):
    if not path.exists():
        return {}
    d = pd.read_csv(path, keep_default_na=False).set_index("question_id")
    out = {}
    for q in d.index:
        out[str(q)] = (str(d.loc[q, lcol]).strip() if lcol in d else "",
                       str(d.loc[q, ccol]).strip().lower() if ccol in d else "",
                       str(d.loc[q, confcol]).strip() if confcol in d else "",
                       str(d.loc[q, notecol]).strip() if notecol in d else "")
    return out


ANNOT = {
    "fable": load_annot(ANA / "gold_commitment_audit_sheet_completed.csv",
                        "manual_gold_label", "manual_gold_commitment",
                        "manual_confidence", "manual_notes"),
    "codex": load_annot(ANA / "gold_commitment_audit_codex.csv",
                        "annotator_gold_label", "annotator_gold_commitment",
                        "annotator_confidence", "annotator_notes"),
    "antigravity": load_annot(ANA / "gold_commitment_audit_antigravity.csv",
                              "annotator_gold_label", "annotator_gold_commitment",
                              "annotator_confidence", "annotator_notes"),
    "human": load_annot(ANA / "gold_commitment_audit_sheet_human.csv",
                        "manual_gold_label", "manual_gold_commitment",
                        "manual_confidence", "manual_notes"),
}


def eff_commit(q, name):
    """Effective commitment: annotator's own; fable falls back to schema."""
    raw = ANNOT[name].get(q, ("", "", "", ""))[1]
    if raw in ("committed", "noncommitted"):
        return raw
    return schema_commit(q) if name == "fable" else ""


# ── Sheet 1: gold + annotators ──────────────────────────────────────────
rows = []
for s in ELIG:
    q = s.question_id
    row = {
        "question_id": q, "lane": s.lane, "answer_type": s.answer_type,
        "schema_confidence": s.schema_confidence, "question": s.question,
        "answer_space": " | ".join(s.answer_space),
        "gold_label": s.gold_label, "schema_commitment": schema_commit(q),
        "gold_label_evidence": s.gold_label_evidence,
        "canonical_claim": s.canonical_claim,
        "ambiguity_notes": s.ambiguity_notes or "",
        "gold_final_answer_text": gold_text.get(q, s.gold_label_evidence),
    }
    for name in ("fable", "codex", "antigravity", "human"):
        lab, com, conf, note = ANNOT[name].get(q, ("", "", "", ""))
        row[f"{name}_label"] = lab
        row[f"{name}_commitment"] = com
        row[f"{name}_conf"] = conf
        row[f"{name}_notes"] = note
    # consensus over the 3 fully-covered annotators (fable falls back to schema)
    ec = [eff_commit(q, n) for n in ("fable", "codex", "antigravity")]
    nc = ec.count("noncommitted")
    cc = ec.count("committed")
    row["majority_commitment"] = ("noncommitted" if nc > cc else
                                  "committed" if cc > nc else schema_commit(q))
    row["commitment_disagreement"] = "YES" if len(set(ec)) > 1 else ""
    rows.append(row)
main = pd.DataFrame(rows)

# ── Sheet 2: model predictions ──────────────────────────────────────────
MODELS = [("ea_full_gemma4", "gemma4"), ("ea_full_qwen3", "qwen3_8b"),
          ("ea_full_hf_qwen36_27b_fp8", "qwen3.6_27b"),
          ("ea_full_hf_gemma_4_31B_it", "gemma4_31b_it"),
          ("gemini_subset16", "gemini_3.1_pro")]
mp = {s.question_id: {"question_id": s.question_id, "lane": s.lane,
                      "gold_label": s.gold_label,
                      "schema_commitment": schema_commit(s.question_id)}
      for s in ELIG}
for run, name in MODELS:
    p = RES / run / "rows.csv"
    if not p.exists():
        continue
    d = pd.read_csv(p)
    d = d[d["round"] == 1]
    for _, r in d.iterrows():
        q = str(r["question_id"])
        if q in mp:
            pred = r["predicted"] if pd.notna(r["predicted"]) else ""
            corr = "" if pd.isna(r["correct"]) else ("Y" if r["correct"] else "N")
            mp[q][f"{name}_pred"] = pred
            mp[q][f"{name}_ok"] = corr
            mp[q][f"{name}_pNC"] = round(float(r["p_noncommit"]), 3) if pd.notna(r.get("p_noncommit")) else ""
models = pd.DataFrame(list(mp.values()))

# ── Sheet 3: disagreements only ─────────────────────────────────────────
dis = main[main["commitment_disagreement"] == "YES"][[
    "question_id", "lane", "answer_type", "question", "gold_label",
    "schema_commitment", "fable_commitment", "codex_commitment",
    "antigravity_commitment", "human_commitment", "majority_commitment",
    "gold_label_evidence"]].copy()

# ── Sheet 4: legend ─────────────────────────────────────────────────────
legend = pd.DataFrame([
    ("question_id / lane / answer_type", "Question identity and schema type"),
    ("answer_space", "The finite label set the question maps to"),
    ("gold_label", "Schema's canonical gold answer label"),
    ("schema_commitment", "committed = gold_label is directional; noncommitted = mixed/conditional/insufficient_data/none_clear"),
    ("gold_label_evidence", "Verbatim quote from the gold response supporting the label"),
    ("canonical_claim / ambiguity_notes", "Schema author's one-line claim and flagged ambiguity"),
    ("gold_final_answer_text", "Full final-answer section of the gold response (source CSV)"),
    ("<annot>_label/_commitment/_conf/_notes", "Each annotator's raw judgement. fable = analyst self-audit (non-blinded; blanks fall back to schema). codex/antigravity = independent BLINDED external raters (139/139). human = partial (~38/139)."),
    ("majority_commitment", "Majority over fable(+schema fallback)/codex/antigravity; tie -> schema"),
    ("commitment_disagreement", "YES if the 3 covered annotators do not all agree on commitment"),
    ("model_predictions sheet", "Each debate model's final-round predicted label, correctness (Y/N), and p_noncommit vs the gold"),
    ("NOTE", "Reanalysis snapshot 2026-07-18. Codex vs Antigravity commit agreement 86.3%, Cohen kappa 0.590; Fleiss 0.531. Hedge-collision conclusion survives all annotator versions."),
], columns=["column", "meaning"])

# ── write + format ──────────────────────────────────────────────────────
WRAP = {"question", "answer_space", "gold_label_evidence", "canonical_claim",
        "ambiguity_notes", "gold_final_answer_text", "fable_notes",
        "codex_notes", "antigravity_notes", "human_notes", "meaning"}
WIDE = {"question": 55, "answer_space": 32, "gold_label_evidence": 60,
        "canonical_claim": 50, "ambiguity_notes": 40,
        "gold_final_answer_text": 90, "fable_notes": 45, "codex_notes": 45,
        "antigravity_notes": 45, "human_notes": 40, "meaning": 90,
        "question_id": 12, "lane": 6, "answer_type": 22}
HEADER = PatternFill("solid", fgColor="1F4E78")
DISFILL = PatternFill("solid", fgColor="FCE4D6")

with pd.ExcelWriter(OUT, engine="openpyxl") as xw:
    main.to_excel(xw, sheet_name="gold_and_annotators", index=False)
    models.to_excel(xw, sheet_name="model_predictions", index=False)
    dis.to_excel(xw, sheet_name="commitment_disagreements", index=False)
    legend.to_excel(xw, sheet_name="legend", index=False)
    for sh in xw.book.worksheets:
        ws = sh
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        # header style
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = HEADER
            c.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
        # column widths + wrap
        headers = [c.value for c in ws[1]]
        for i, h in enumerate(headers, 1):
            L = get_column_letter(i)
            ws.column_dimensions[L].width = WIDE.get(h, 15)
            if h in WRAP:
                for cell in ws[L][1:]:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[1].height = 30
    # highlight disagreement rows on main sheet
    ws = xw.book["gold_and_annotators"]
    dcol = [c.value for c in ws[1]].index("commitment_disagreement") + 1
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, dcol).value == "YES":
            for c in range(1, ws.max_column + 1):
                ws.cell(r, c).fill = DISFILL

print(f"wrote {OUT}")
print(f"  gold_and_annotators: {len(main)} rows x {len(main.columns)} cols; "
      f"disagreements: {len(dis)}")
print(f"  model_predictions: {len(models)} rows ({[m[1] for m in MODELS]})")
