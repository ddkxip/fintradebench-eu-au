"""Read the workbook. Only allowlisted fields enter model prompts."""

import math
import re
from pathlib import Path

from openpyxl import load_workbook

DEFAULT_DATA = Path(__file__).parent / "sp500_valuation_technical_2026-04-27_2.xlsx"
FIELDS = (
    "PE_Ratio", "PB_Ratio", "PS_Ratio", "EV_EBITDA", "RSI_14", "SMA_50",
    "Price_to_SMA50", "Volume_to_AvgVol20", "Close_Price", "Volume",
    "Fundamental_Quarter_Date",
)


def load_companies(path=DEFAULT_DATA, rows=5, snapshot_date="2026-04-27",
                   identity_mode="fake_ticker"):
    if identity_mode not in ("fake_ticker", "real_name"):
        raise ValueError("identity_mode must be fake_ticker or real_name")
    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        source = list(workbook.active.iter_rows(values_only=True))
    finally:
        workbook.close()
    headers = source[1]  # Workbook title is row 1; headers are row 2.
    required = {*FIELDS, "Ticker", "Company_Name", "Fake_Ticker"}
    if not required.issubset(headers):
        raise ValueError(f"Missing columns: {required - set(headers)}")
    records = [dict(zip(headers, values)) for values in source[2:] if any(v is not None for v in values)]
    if not 1 <= rows <= len(records):
        raise ValueError(f"--rows must be between 1 and {len(records)}")
    aliases = [record["Fake_Ticker"] for record in records]
    if len(set(aliases)) != len(aliases) or any(not isinstance(t, str) or not re.fullmatch(r"[A-Z]{4}", t) for t in aliases):
        raise ValueError("Fake tickers must be unique four-letter uppercase names")
    companies = []
    for row_number, record in enumerate(records[:rows], 3):
        evidence = {"FakeTicker": record["Fake_Ticker"], "SnapshotDate": snapshot_date}
        if identity_mode == "real_name":
            name = record["Company_Name"]
            if not isinstance(name, str) or not name.strip():
                raise ValueError(f"Row {row_number}: missing company name")
            evidence["CompanyName"] = name
        evidence.update({key: record[key] for key in FIELDS})
        for key in FIELDS[:-1]:
            value = evidence[key]
            if value is not None and (not isinstance(value, (int, float)) or not math.isfinite(value)):
                raise ValueError(f"Row {row_number}: invalid number in {key}")
        ratio = evidence["Price_to_SMA50"]
        evidence["PriceVsSMA50_Percent"] = round(100 * (ratio - 1), 2) if ratio is not None else None
        evidence["Sector_PE"] = None
        companies.append({"excel_row": row_number, "evidence": evidence})
    # This full identity list is for offline checks only.
    identities = [str(record[key]) for record in records for key in ("Ticker", "Company_Name") if record[key]]
    return companies, identities
